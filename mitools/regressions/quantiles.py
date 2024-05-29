import hashlib
import pickle
import re
import traceback
import warnings
from dataclasses import dataclass
from os import PathLike
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Union

import numpy as np
import openpyxl
import pandas as pd
import statsmodels.formula.api as smf
from numpy.linalg import LinAlgError
from pandas import DataFrame, MultiIndex, Series
from statsmodels.regression.linear_model import RegressionResultsWrapper
from tqdm.notebook import tqdm

from mitools.economic_complexity import StringMapper
from mitools.pandas import idxslice
from mitools.regressions import generate_hash_from_dataframe
from mitools.utils import auto_adjust_columns_width

warnings.simplefilter("ignore")
Color = Union[Tuple[int, int, int], str]


@dataclass(frozen=True)
class QuantileRegStrs:
    UNNAMED: str = "Unnamed: 0"
    COEF: str = "coef"
    T_VALUE: str = "t"
    P_VALUE: str = "P>|t|"
    VALUE: str = "Value"
    QUANTILE: str = "Quantile"
    INDEPENDENT_VARS: str = "Independent Vars"
    REGRESSION_TYPE: str = "Regression Type"
    REGRESSION_DEGREE: str = "Regression Degree"
    DEPENDENT_VAR: str = "Dependent Var"
    VARIABLE_TYPE: str = "Variable Type"
    EXOG_VAR: str = "Exog"
    CONTROL_VAR: str = "Control"
    ID: str = "Id"
    QUADRATIC_REG: str = "quadratic"
    LINEAR_REG: str = "linear"
    QUADRATIC_VAR_SUFFIX: str = "_square"
    INDEPENDENT_VARS_PATTERN: str = r"^I\((.*)\)$"
    STATS: str = "Stats"
    INTERCEPT: str = "Intercept"
    ANNOTATION: str = "Q"
    PARQUET_SUFFIX: str = "regressions"
    EXCEL_SUFFIX: str = "regressions"
    MAIN_PLOT: str = "regression_data"
    PLOTS_SUFFIX: str = "regression"
    ADJ_METHOD: str = "Adj Method"
    DATE: str = "Date"
    TIME: str = "Time"
    PSEUDO_R_SQUARED: str = "Pseudo R-squared"
    BANDWIDTH: str = "Bandwidth"
    SPARSITY: str = "Sparsity"
    N_OBSERVATIONS: str = "N Observations"
    DF_RESIDUALS: str = "Df Residuals"
    DF_MODEL: str = "Df Model"
    KURTOSIS: str = "Kurtosis"
    SKEWNESS: str = "Skewness"


class QuantilesRegressionSpecs:
    def __init__(
        self,
        dependent_variable: str,
        independent_variables: List[str],
        quantiles: List[float],
        quadratic: bool,
        regression_type: str,
        data: DataFrame,
        group: Optional[str] = None,
        control_variables: Optional[List[str]] = None,
        panel: Optional[bool] = False,
    ):
        self.dependent_variable = dependent_variable
        self.independent_variables = independent_variables
        self.quadratic = quadratic
        if self.quadratic and not any(
            [
                f"{var}{QuantileRegStrs.QUADRATIC_VAR_SUFFIX}"
                in self.independent_variables
                for var in self.independent_variables
            ]
        ):
            self.independent_variables += [
                f"{var}{QuantileRegStrs.QUADRATIC_VAR_SUFFIX}"
                for var in independent_variables
            ]
        self.independent_variables.sort()
        self.control_variables = control_variables or []
        self.control_variables.sort()
        self.variables = (
            [self.dependent_variable]
            + self.independent_variables
            + self.control_variables
        )
        self.quantiles = quantiles
        self.regression_type = regression_type
        self.data = data
        self.regression_id = create_regression_id(
            self.regression_type,
            self.quadratic,
            self.dependent_variable,
            self.independent_variables,
            self.control_variables,
        )
        self.group = group
        self.formula = self.get_formula()

    def get_formula(self, str_mapper: Optional[StringMapper] = None) -> str:
        if str_mapper:
            independent_variables = str_mapper.prettify_strs(self.independent_variables)
            control_variables = str_mapper.prettify_strs(self.control_variables)
            dependent_variable = str_mapper.prettify_str(self.dependent_variable)
        else:
            independent_variables = self.independent_variables
            control_variables = self.control_variables
            dependent_variable = self.dependent_variable
        formula_terms = [
            var
            for var in independent_variables
            if QuantileRegStrs.QUADRATIC_VAR_SUFFIX not in var
        ]
        formula_terms += [
            f"I({var})"
            for var in independent_variables
            if QuantileRegStrs.QUADRATIC_VAR_SUFFIX in var
        ]
        if control_variables:
            formula_terms += control_variables
        formula = f"{dependent_variable} ~ " + " + ".join(formula_terms)
        return formula

    def data_statistics_table(self, str_mapper: Optional[StringMapper] = None):
        table = self.data[[self.variables]].describe(percentiles=[0.5]).T
        table.columns = [
            QuantileRegStrs.N_OBSERVATIONS,
            "Mean",
            "Std. Dev.",
            "Min",
            "Median",
            "Max",
        ]
        table[QuantileRegStrs.KURTOSIS] = self.data[[self.variables]].kurtosis()
        table[QuantileRegStrs.SKEWNESS] = self.data[[self.variables]].skew()
        table[QuantileRegStrs.N_OBSERVATIONS] = table[
            QuantileRegStrs.N_OBSERVATIONS
        ].astype(int)
        numeric_cols = [c for c in table.columns if c != QuantileRegStrs.N_OBSERVATIONS]
        table[numeric_cols] = table[numeric_cols].round(7)
        table.columns = (
            pd.MultiIndex.from_product([[self.group], table.columns])
            if self.group
            else table.columns
        )
        if str_mapper:
            table.index = table.index.map(lambda x: str_mapper.prettify_str(x))
        return table.sort_index(ascending=True)

    def data_statistics_latex_table(self, str_mapper: Optional[StringMapper] = None):
        table = self.data_statistics_table(str_mapper)
        symbols_pattern = r"([\ \_\-\&\%\$\#])"
        table = table.rename(
            index=lambda x: re.sub(symbols_pattern, regex_symbol_replacement, x)
            if isinstance(x, str)
            else str(round(x, 1))
        )
        table_latex = table.to_latex(
            multirow=True, multicolumn=True, multicolumn_format="c"
        )
        table_text = (
            "\\begin{adjustbox}{width=\\textwidth,center}\n"
            + f"{table_latex}"
            + "\end{adjustbox}\n"
        )
        return table_text

    def store(self, folder_path: Path):
        self.data = None
        with open(folder_path / f"{self.regression_id}.reg_specs", "wb") as file:
            pickle.dump(self, file)


class QuantilesRegression:
    def __init__(self, coeffs, stats):
        self.coeffs = coeffs
        self.stats = stats

        self.id = self.coeffs.index.get_level_values(QuantileRegStrs.ID).tolist()[0]
        self.group = self.coeffs.columns.tolist()[0]

        self.dependent_variables = self.coeffs.index.get_level_values(
            QuantileRegStrs.DEPENDENT_VAR
        ).tolist()[0]

        self.independent_variables = (
            self.coeffs.loc[
                self.coeffs.index.get_level_values(QuantileRegStrs.VARIABLE_TYPE)
                == QuantileRegStrs.EXOG_VAR
            ]
            .index.get_level_values(QuantileRegStrs.INDEPENDENT_VARS)
            .unique()
            .tolist()
        )
        self.control_variables = (
            self.coeffs.loc[
                self.coeffs.index.get_level_values(QuantileRegStrs.VARIABLE_TYPE)
                == QuantileRegStrs.CONTROL_VAR
            ]
            .index.get_level_values(QuantileRegStrs.INDEPENDENT_VARS)
            .unique()
            .tolist()
        )

        self.quantiles = (
            self.coeffs.index.get_level_values(QuantileRegStrs.QUANTILE)
            .unique()
            .tolist()
        )
        self.quadratic = (
            self.coeffs.index.get_level_values(
                QuantileRegStrs.REGRESSION_DEGREE
            ).tolist()[0]
            == QuantileRegStrs.QUADRATIC_REG
        )
        self.regression_type = self.coeffs.index.get_level_values(
            QuantileRegStrs.REGRESSION_TYPE
        ).tolist()[0]

    def coefficients(self, quantiles: Optional[List[float]] = None):
        if quantiles is None:
            return self.coeffs
        return self.coeffs.loc[
            self.coeffs.index.get_level_values(QuantileRegStrs.QUANTILE).isin(quantiles)
        ]

    def n_obs(self, quantiles: Optional[List[float]] = None):
        if quantiles is None:
            stats = self.stats.loc[(slice(None), QuantileRegStrs.N_OBSERVATIONS), :]
        else:
            stats = self.stats.loc[(quantiles, QuantileRegStrs.N_OBSERVATIONS), :]
        stats.index = stats.index.droplevel(QuantileRegStrs.STATS)
        stats.columns = [QuantileRegStrs.N_OBSERVATIONS]
        return stats

    def r_squared(self, quantiles: Optional[List[float]] = None):
        if quantiles is None:
            stats = self.stats.loc[(slice(None), QuantileRegStrs.PSEUDO_R_SQUARED), :]
        else:
            stats = self.stats.loc[(quantiles, QuantileRegStrs.PSEUDO_R_SQUARED), :]
        stats.index = stats.index.droplevel(QuantileRegStrs.STATS)
        stats.columns = [QuantileRegStrs.PSEUDO_R_SQUARED]
        return stats

    def coefficients_quantiles_table(self, quantiles: Optional[List[float]] = None):
        table = self.coeffs.unstack(level=QuantileRegStrs.QUANTILE)
        if quantiles is not None:
            table = table.loc[:, (slice(None), quantiles)]
        return table.sort_index(
            axis=0,
            level=[QuantileRegStrs.VARIABLE_TYPE, QuantileRegStrs.INDEPENDENT_VARS],
            ascending=[False, True],
        )

    def coefficients_quantiles_latex_table(
        self,
        quantiles: Optional[List[float]] = None,
        note: Optional[bool] = False,
        str_mapper: Optional[StringMapper] = None,
    ):
        table = self.coefficients_quantiles_table(quantiles).droplevel(
            [
                QuantileRegStrs.ID,
                QuantileRegStrs.REGRESSION_TYPE,
                QuantileRegStrs.REGRESSION_DEGREE,
                QuantileRegStrs.VARIABLE_TYPE,
            ],
            axis=0,
        )
        if str_mapper is not None:
            levels_to_remap = [
                QuantileRegStrs.DEPENDENT_VAR,
                QuantileRegStrs.INDEPENDENT_VARS,
            ]
            pretty_index = table.index.set_levels(
                [
                    prettify_index_level(
                        str_mapper,
                        QuantileRegStrs.QUADRATIC_VAR_SUFFIX,
                        level,
                        level_id,
                        levels_to_remap,
                    )
                    for level, level_id in zip(table.index.levels, table.index.names)
                ],
                level=table.index.names,
            )
            table.index = pretty_index
        symbols_pattern = r"([\ \_\-\&\%\$\#])"
        table = table.rename(
            columns=lambda x: re.sub(symbols_pattern, regex_symbol_replacement, x)
            if isinstance(x, str)
            else str(round(x, 1)),
            index=lambda x: re.sub(symbols_pattern, regex_symbol_replacement, x)
            if isinstance(x, str)
            else str(round(x, 1)),
        ).to_latex(multirow=True, multicolumn=True, multicolumn_format="c")
        table_text = (
            "\\begin{adjustbox}{width=\\textwidth,center}\n"
            + f"{table}"
            + "\end{adjustbox}\n"
        )
        table_text = (
            table_text
            + "{\\centering\\tiny Note: * p\\textless0.05, ** p\\textless0.01, *** p\\textless0.001\\par}"
            if note
            else table_text
        )
        print(table_text)

    def model_specification(self, str_mapper: Optional[StringMapper] = None):
        if str_mapper:
            independent_variables = [
                str_mapper.prettify_str(var)
                if QuantileRegStrs.QUADRATIC_VAR_SUFFIX not in var
                else f"{str_mapper.prettify_str(var.replace(QuantileRegStrs.QUADRATIC_VAR_SUFFIX, ''))}{QuantileRegStrs.QUADRATIC_VAR_SUFFIX}"
                for var in self.independent_variables
            ]
            control_variables = [
                str_mapper.prettify_str(var) for var in self.control_variables
            ]
        else:
            independent_variables = self.independent_variables
            control_variables = self.control_variables
        model_specification = f"{self.dependent_variables if not str_mapper else str_mapper.prettify_str(self.dependent_variables)}"
        model_specification += f" ~ {' + '.join(independent_variables)}"
        model_specification += (
            f" + {' + '.join([var for var in control_variables if var != 'Intercept'])}"
            if control_variables
            else ""
        )
        model_specification = model_specification.split(" + ")
        lines = []
        line = ""
        for string in model_specification[:-1]:
            if len(line) + len(string) < 120:
                line += f"{string} + "
            else:
                lines.append(line + r"\\")
                line = string + " + "
        lines.append(model_specification[-1])
        model_specification = "".join(lines)
        symbols_pattern = r"([\ \_\-\&\%\$\#])"
        model_specification = re.sub(
            symbols_pattern, regex_symbol_replacement, model_specification
        ).replace("~", "\\sim")
        print(f"${model_specification}$")

    def abstract_model_specification(self):
        pass

    def quantile_model_equation(self):
        print(
            "$\\min_{\\beta} \\sum_{i:y_g \\geq x_g^T\\beta} q |y_g - x_g^T\\beta| + \\sum_{g:y_g < x_g^T\\beta} (1-q) |y_g - x_g^T\\beta|$"
        )

    def store(self, folder_path: Path):
        with open(folder_path / f"{self.id}.reg_coeffs", "wb") as file:
            pickle.dump(self, file)


def create_regression_id(
    regression_type: str,
    regression_degree: str,
    regression_dependent_var: str,
    regression_indep_vars: List[str],
    control_variables: List[str],
    id_len: Optional[int] = 6,
) -> str:
    str_to_hash = " ".join(
        [
            regression_type,
            regression_degree,
        ]
    )
    id_hasher = hashlib.md5()
    id_hasher.update(rf"{str_to_hash}".encode("utf-8"))
    kind_id = id_hasher.hexdigest()[:id_len]

    id_hasher = hashlib.md5()
    id_hasher.update(rf"{regression_dependent_var}".encode("utf-8"))
    dep_id = id_hasher.hexdigest()[:id_len]

    str_to_hash = " ".join([v for v in regression_indep_vars if "_square" not in v])
    id_hasher = hashlib.md5()
    id_hasher.update(rf"{str_to_hash}".encode("utf-8"))
    indep_id = id_hasher.hexdigest()[:id_len]

    control_vars_str = " ".join([v for v in control_variables])
    id_hasher = hashlib.md5()
    id_hasher.update(rf"{control_vars_str}".encode("utf-8"))
    control_vars_id = id_hasher.hexdigest()[:id_len] if control_variables else "None"
    return f"{kind_id}-{dep_id}-{indep_id}-{control_vars_id}"


def create_regression_file_paths(
    eci_type_folder: PathLike, regression_id: str
) -> Tuple[Path, Path]:
    eci_type_folder = Path(eci_type_folder)
    main_plot = eci_type_folder / f"{QuantileRegStrs.MAIN_PLOT}.png"
    regression_plot = (
        eci_type_folder / f"{regression_id}_{QuantileRegStrs.PLOTS_SUFFIX}.png"
    )
    return main_plot, regression_plot


def get_regression_predictions(
    data: DataFrame,
    regression_coeffs: DataFrame,
    groups: List[str],
    groups_col: str,
    all_groups: str,
) -> Tuple[DataFrame, DataFrame, DataFrame]:
    predictions, significances, x_values = [], [], []
    for group in groups:
        group_predictions, group_significances, group_x_values = (
            get_quantile_regression_predictions_by_group(
                regression_data=data,
                regression_coeffs=regression_coeffs,
                group=group,
                groups_col=groups_col,
                all_groups=all_groups,
            )
        )
        group_predictions.columns = pd.MultiIndex.from_tuples(
            [(group,) + c for c in group_predictions.columns]
        )
        group_significances.columns = pd.MultiIndex.from_tuples(
            [(group,) + c for c in group_significances.columns]
        )
        group_x_values.columns = pd.MultiIndex.from_tuples(
            [(group, c) for c in group_x_values.columns]
        )
        predictions.append(group_predictions)
        significances.append(group_significances)
        x_values.append(group_x_values)
    predictions = pd.concat(predictions, axis=1)
    significances = pd.concat(significances, axis=1)
    x_values = pd.concat(x_values, axis=1)
    return predictions, significances, x_values


def prettify_index_level(
    mapper: StringMapper,
    pattern: str,
    level: str,
    level_name: str,
    levels_to_remap: List[str],
) -> str:
    if level_name in levels_to_remap:
        return level.map(lambda x: prettify_with_pattern(x, mapper, pattern))
    return level


def prettify_with_pattern(string: str, mapper: StringMapper, pattern: str) -> str:
    base_string, pattern_str, _ = string.partition(pattern)
    remapped_base = mapper.prettify_str(base_string)
    return f"{remapped_base}{pattern}" if pattern_str else remapped_base


def get_quantile_regression_predictions_by_group(
    regression_data: DataFrame,
    regression_coeffs: DataFrame,
    group: str,
    groups_col: str,
    all_groups: str,
) -> Tuple[DataFrame, DataFrame]:
    group_data = get_group_data(regression_data, group, groups_col, all_groups)
    dependent_var = regression_coeffs.index.get_level_values(
        QuantileRegStrs.DEPENDENT_VAR
    ).unique()[0]
    quantiles = regression_coeffs.index.get_level_values(
        QuantileRegStrs.QUANTILE
    ).unique()
    independent_vars = [
        var
        for var in regression_coeffs.index.get_level_values(
            QuantileRegStrs.INDEPENDENT_VARS
        ).unique()
        if var != QuantileRegStrs.INTERCEPT
    ]
    x_values = prepare_x_values(group_data, independent_vars)
    predictions, significances, columns = [], [], []
    for var in independent_vars:
        quadratic = QuantileRegStrs.QUADRATIC_VAR_SUFFIX in var
        x_var_values = x_values[var.replace(QuantileRegStrs.QUADRATIC_VAR_SUFFIX, "")]
        var_values = (
            [var, QuantileRegStrs.INTERCEPT]
            if not quadratic
            else [
                var.replace(QuantileRegStrs.QUADRATIC_VAR_SUFFIX, ""),
                var,
                QuantileRegStrs.INTERCEPT,
            ]
        )
        vars_idx = idxslice(
            regression_coeffs,
            level=QuantileRegStrs.INDEPENDENT_VARS,
            value=var_values,
            axis=0,
        )
        var_coeffs = regression_coeffs.loc[vars_idx, :]
        for quantile in quantiles:
            quantile_idx = idxslice(
                var_coeffs, level=QuantileRegStrs.QUANTILE, value=quantile, axis=0
            )
            values = var_coeffs.loc[quantile_idx, group].values
            significance = ",".join(
                [
                    match.group() if match else "-"
                    for match in [re.search(r"\*+$", val) for val in values[:-1]]
                ]
            )
            coeffs = [
                float(re.search(r"([-\d.eE]+)\(", val).group(1)) for val in values
            ]
            coeffs_names = (
                [
                    QuantileRegStrs.LINEAR_REG,
                    QuantileRegStrs.QUADRATIC_REG,
                    QuantileRegStrs.INTERCEPT,
                ]
                if quadratic
                else [QuantileRegStrs.LINEAR_REG, QuantileRegStrs.INTERCEPT]
            )
            coeffs = dict(zip(coeffs_names, coeffs))
            prediction = get_prediction(x_var_values, coeffs, quadratic)

            predictions.append(prediction)
            significances.append(significance)
            columns.append((dependent_var, var, quantile))

    predictions = DataFrame(predictions, index=MultiIndex.from_tuples(columns)).T
    significances = DataFrame(significances, index=MultiIndex.from_tuples(columns)).T
    return predictions, significances, x_values


def get_group_data(
    regression_data: DataFrame, group: str, group_col: str, all_groups: str
) -> DataFrame:
    if group != all_groups:
        return regression_data.loc[
            regression_data.index.get_level_values(group_col) == group
        ]
    return regression_data


def prepare_x_values(group_data: DataFrame, independent_vars: List[str]) -> DataFrame:
    return DataFrame(
        {
            var: np.linspace(group_data[var].min(), group_data[var].max(), 100)
            for var in independent_vars
            if QuantileRegStrs.QUADRATIC_VAR_SUFFIX not in var
        }
    )


def get_prediction(x_values: Series, coeffs: List[float], quadratic: bool) -> float:
    prediction = (
        coeffs[QuantileRegStrs.INTERCEPT]
        + coeffs[QuantileRegStrs.LINEAR_REG] * x_values
    )
    if quadratic:
        prediction += coeffs[QuantileRegStrs.QUADRATIC_REG] * x_values**2
    return prediction


def get_quantile_regression_results_stats(
    results: Dict[int, RegressionResultsWrapper],
) -> DataFrame:
    _stats_name_remap_dict = {
        "Dep. Variable:": QuantileRegStrs.DEPENDENT_VAR,
        "Model:": QuantileRegStrs.REGRESSION_TYPE,
        "Method:": QuantileRegStrs.ADJ_METHOD,
        "Date:": QuantileRegStrs.DATE,
        "Time:": QuantileRegStrs.TIME,
        "Pseudo R-squared:": QuantileRegStrs.PSEUDO_R_SQUARED,
        "Bandwidth:": QuantileRegStrs.BANDWIDTH,
        "Sparsity:": QuantileRegStrs.SPARSITY,
        "No. Observations:": QuantileRegStrs.N_OBSERVATIONS,
        "Df Residuals:": QuantileRegStrs.DF_RESIDUALS,
        "Df Model:": QuantileRegStrs.DF_MODEL,
    }
    regression_stats = []
    for q, result in results.items():
        stats = result.summary().tables[0].as_html()
        stats = pd.read_html(stats, index_col=0)[0].reset_index()
        stats = pd.concat(
            [stats.iloc[:-1, :2], stats.iloc[:, 2:].rename(columns={2: 0, 3: 1})],
            axis=0,
            ignore_index=True,
        )
        stats.columns = [QuantileRegStrs.STATS, QuantileRegStrs.VALUE]
        stats[QuantileRegStrs.QUANTILE] = q
        stats = stats.set_index([QuantileRegStrs.QUANTILE, QuantileRegStrs.STATS])
        regression_stats.append(stats)
    regression_stats = pd.concat(regression_stats, axis=0)
    regression_stats.index = regression_stats.index.set_levels(
        regression_stats.index.levels[
            regression_stats.index.names.index(QuantileRegStrs.STATS)
        ].map(_stats_name_remap_dict.get),
        level=QuantileRegStrs.STATS,
    )
    return regression_stats


def process_result_wrappers_coeffs(
    results: Dict[int, RegressionResultsWrapper],
) -> DataFrame:
    return pd.concat(
        [
            process_quantile_regression_coeffs_result(q, result)
            for q, result in results.items()
        ],
        axis=1,
    )


def melt_and_rename_regression_coeffs(regression_coeffs: DataFrame) -> DataFrame:
    melted_df = regression_coeffs.reset_index().melt(
        id_vars=QuantileRegStrs.UNNAMED,
        var_name=[QuantileRegStrs.QUANTILE],
        value_name=QuantileRegStrs.VALUE,
    )
    melted_df.columns = [QuantileRegStrs.INDEPENDENT_VARS, *melted_df.columns[1:]]
    return melted_df


def update_regression_coeffs_independent_vars(
    regression_coeffs: DataFrame,
) -> DataFrame:
    regression_coeffs[QuantileRegStrs.INDEPENDENT_VARS] = regression_coeffs[
        QuantileRegStrs.INDEPENDENT_VARS
    ].replace(QuantileRegStrs.INDEPENDENT_VARS_PATTERN, r"\1", regex=True)
    return regression_coeffs


def set_regression_coeffs_info(
    regression_coeffs: DataFrame,
    results: Dict[int, RegressionResultsWrapper],
    independent_variables: List[str],
) -> DataFrame:
    regression_coeffs[QuantileRegStrs.REGRESSION_TYPE] = type(
        list(results.values())[0].model
    ).__name__
    reg_degree = (
        QuantileRegStrs.QUADRATIC_REG
        if all(
            f"{var}{QuantileRegStrs.QUADRATIC_VAR_SUFFIX}"
            in regression_coeffs[QuantileRegStrs.INDEPENDENT_VARS].values
            for var in independent_variables
        )
        else QuantileRegStrs.LINEAR_REG
    )
    regression_coeffs[QuantileRegStrs.REGRESSION_DEGREE] = reg_degree
    regression_coeffs[QuantileRegStrs.DEPENDENT_VAR] = list(results.values())[
        0
    ].model.endog_names
    return regression_coeffs


def classify_regression_coeffs_variables(
    regression_coeffs: DataFrame, independent_variables: List[str]
) -> DataFrame:
    regression_coeffs[QuantileRegStrs.VARIABLE_TYPE] = regression_coeffs[
        QuantileRegStrs.INDEPENDENT_VARS
    ].apply(
        lambda x: QuantileRegStrs.EXOG_VAR
        if x.replace(QuantileRegStrs.QUADRATIC_VAR_SUFFIX, "") in independent_variables
        else QuantileRegStrs.CONTROL_VAR
    )
    return regression_coeffs


def sort_and_set_regression_coeffs_index(regression_coeffs: DataFrame) -> DataFrame:
    regression_coeffs = regression_coeffs.sort_values(
        by=[
            QuantileRegStrs.VARIABLE_TYPE,
            QuantileRegStrs.INDEPENDENT_VARS,
            QuantileRegStrs.QUANTILE,
        ],
        ascending=[False, True, True],
    )
    regression_coeffs[QuantileRegStrs.QUANTILE] = regression_coeffs[
        QuantileRegStrs.QUANTILE
    ].astype(float)
    regression_coeffs = regression_coeffs.set_index(
        [
            QuantileRegStrs.REGRESSION_TYPE,
            QuantileRegStrs.REGRESSION_DEGREE,
            QuantileRegStrs.DEPENDENT_VAR,
            QuantileRegStrs.VARIABLE_TYPE,
            QuantileRegStrs.INDEPENDENT_VARS,
            QuantileRegStrs.QUANTILE,
        ]
    )
    return regression_coeffs


def add_id_and_reorder_regression_coeffs(regression_coeffs: DataFrame) -> DataFrame:
    regression_coeffs[QuantileRegStrs.ID] = generate_hash_from_dataframe(
        regression_coeffs,
        [
            QuantileRegStrs.REGRESSION_TYPE,
            QuantileRegStrs.REGRESSION_DEGREE,
            QuantileRegStrs.INDEPENDENT_VARS,
        ],
        length=12,
    )
    regression_coeffs = regression_coeffs.set_index(QuantileRegStrs.ID, append=True)
    regression_coeffs = regression_coeffs.reorder_levels(
        [regression_coeffs.index.names[-1]] + regression_coeffs.index.names[:-1]
    )
    return regression_coeffs


def process_quantile_regression_coeffs_result(
    q: float, result: RegressionResultsWrapper
) -> DataFrame:
    coeffs = pd.concat(pd.read_html(result.summary().tables[1].as_html(), header=0))
    coeffs = coeffs.set_index(QuantileRegStrs.UNNAMED)
    coeffs[QuantileRegStrs.VALUE] = coeffs[
        [QuantileRegStrs.COEF, QuantileRegStrs.T_VALUE, QuantileRegStrs.P_VALUE]
    ].apply(quantile_regression_value, axis=1)
    coeffs = coeffs[[QuantileRegStrs.VALUE]]
    coeffs.columns = pd.MultiIndex.from_tuples([(str(q), c) for c in coeffs.columns])
    return coeffs


def get_quantile_regression_results_coeffs(
    results: Dict[int, RegressionResultsWrapper], independent_variables: List[str]
) -> DataFrame:
    regression_coeffs = process_result_wrappers_coeffs(results)
    regression_coeffs = melt_and_rename_regression_coeffs(regression_coeffs)
    regression_coeffs = update_regression_coeffs_independent_vars(regression_coeffs)
    regression_coeffs = set_regression_coeffs_info(
        regression_coeffs, results, independent_variables
    )
    regression_coeffs = classify_regression_coeffs_variables(
        regression_coeffs, independent_variables
    )
    regression_coeffs = sort_and_set_regression_coeffs_index(regression_coeffs)
    # regression_coeffs = add_id_and_reorder_regression_coeffs(regression_coeffs)
    return regression_coeffs


def quantile_regression_value(row: Series) -> Series:
    coeff = round(row[QuantileRegStrs.COEF], 5)
    t_value = round(row[QuantileRegStrs.T_VALUE], 5)
    p_value = row[QuantileRegStrs.P_VALUE]
    if p_value <= 0.001:
        return f"{coeff}({t_value})***"
    elif p_value <= 0.01:
        return f"{coeff}({t_value})**"
    elif p_value <= 0.05:
        return f"{coeff}({t_value})*"
    else:
        return f"{coeff}({t_value})"


def get_quantile_regression_results(
    regression: "QuantilesRegressionSpecs", max_iter: Optional[int] = 2_500
) -> Dict[float, RegressionResultsWrapper]:
    results = {
        q: smf.quantreg(regression.formula, regression.data).fit(q=q, max_iter=max_iter)
        for q in regression.quantiles
    }
    return results


def prepare_regression_data(
    data: DataFrame,
    dependent_variable: str,
    independent_variables: List[str],
    control_variables: Optional[List[str]] = None,
    str_mapper: Optional[StringMapper] = None,
) -> Tuple[DataFrame, str, List[str], List[str]]:
    if control_variables is None:
        control_variables = []
    regression_data = data.loc[
        :, [dependent_variable, *independent_variables, *control_variables]
    ].copy(deep=True)
    if str_mapper is not None:
        regression_data.columns = [
            str_mapper.uglify_str(var) for var in regression_data.columns
        ]
        dependent_variable = str_mapper.uglify_str(dependent_variable)
        independent_variables = [
            str_mapper.uglify_str(var) for var in independent_variables
        ]
        control_variables = [str_mapper.uglify_str(var) for var in control_variables]
    return regression_data, dependent_variable, independent_variables, control_variables


def create_quantile_regressions_results(
    data: DataFrame,
    dependent_variables: List[str],
    independent_variables: Dict[str, List[str]],
    groups: List[str],
    regressions_folder: PathLike,
    quadratics: List[bool],
    quantiles: Optional[List[float]] = None,
    str_mapper: Optional[StringMapper] = None,
    control_variables: Optional[List[List[str]]] = None,
    group_col: Optional[str] = "Income Group",
    all_groups: Optional[str] = "All income",
    max_iter: Optional[int] = 2_500,
    recalculate: Optional[bool] = False,
):
    regressions_info = {}
    regressions = {}
    if control_variables is None:
        control_variables = [[]]
    for dependent_variable in tqdm(
        dependent_variables, desc="Dependent Variables", position=0, leave=False
    ):
        regressions_info[dependent_variable] = {}
        regressions[dependent_variable] = {}
        dep_var_name = dependent_variable.replace("/", "").replace(" ", "_")
        dep_var_folder = regressions_folder / dep_var_name
        if not dep_var_folder.exists():
            dep_var_folder.mkdir(exist_ok=True)
        for name_tag, indep_variables in tqdm(
            independent_variables.items(),
            desc="Independent Variables",
            position=1,
            leave=False,
        ):
            name_tag_folder = dep_var_folder / name_tag
            if not name_tag_folder.exists():
                name_tag_folder.mkdir(exist_ok=True)
            dep_var_name_excel = (
                name_tag_folder
                / f"{name_tag}_{dep_var_name}_{QuantileRegStrs.EXCEL_SUFFIX}.xlsx"
            )
            dep_var_name_parquet = (
                name_tag_folder
                / f"{name_tag}_{dep_var_name}_{QuantileRegStrs.PARQUET_SUFFIX}.parquet"
            )
            if (
                not dep_var_name_excel.exists()
                or not dep_var_name_parquet.exists()
                or recalculate
            ):
                name_tag_regressions = []
                for control_vars in tqdm(
                    control_variables, desc="Control Variables", position=2, leave=False
                ):
                    if dependent_variable not in control_vars:
                        try:
                            quadratic_regressions = []
                            for quadratic in quadratics:
                                group_regressions = []
                                for group in groups:
                                    if group != all_groups:
                                        group_data = data.loc[
                                            idxslice(
                                                data,
                                                level=group_col,
                                                value=group,
                                                axis=0,
                                            ),
                                            :,
                                        ].copy(deep=True)
                                    else:
                                        group_data = data.copy(deep=True)
                                    (
                                        regression_data,
                                        dependent_var,
                                        independent_vars,
                                        c_vars,
                                    ) = prepare_regression_data(
                                        data=group_data,
                                        dependent_variable=dependent_variable,
                                        independent_variables=indep_variables,
                                        control_variables=control_vars,
                                        str_mapper=str_mapper,
                                    )
                                    regression_info = QuantilesRegressionSpecs(
                                        group=group,
                                        dependent_variable=dependent_var,
                                        independent_variables=independent_vars,
                                        control_variables=c_vars,
                                        quantiles=quantiles,
                                        quadratic="quadratic"
                                        if quadratic
                                        else "linear",
                                        data=regression_data,
                                        regression_type="ols",
                                    )
                                    regression_id = regression_info.regression_id
                                    with warnings.catch_warnings():
                                        regression_results = (
                                            get_quantile_regression_results(
                                                regression=regression_info,
                                                max_iter=max_iter,
                                            )
                                        )

                                    regression_coeffs = (
                                        get_quantile_regression_results_coeffs(
                                            results=regression_results,
                                            independent_variables=independent_vars,
                                        )
                                    )
                                    regression_coeffs.columns = [group]
                                    regression_coeffs["Id"] = (
                                        regression_info.regression_id
                                    )
                                    regression_coeffs = regression_coeffs.set_index(
                                        "Id", append=True
                                    )
                                    regression_coeffs = (
                                        regression_coeffs.reorder_levels(
                                            [-1, 0, 1, 2, 3, 4, 5]
                                        )
                                    )
                                    regression_stats = (
                                        get_quantile_regression_results_stats(
                                            results=regression_results
                                        )
                                    )

                                    regression = QuantilesRegression(
                                        coeffs=regression_coeffs, stats=regression_stats
                                    )
                                    regressions[dependent_variable].setdefault(
                                        regression.id, []
                                    )
                                    regressions[dependent_variable][
                                        regression.id
                                    ].append(regression)

                                    regressions_info[dependent_variable].setdefault(
                                        regression.id, []
                                    )
                                    regressions_info[dependent_variable][
                                        regression.id
                                    ].append(regression_info)
                                    regression_info.store(name_tag_folder)
                                    group_regressions.append(regression.coeffs)
                                group_regressions = pd.concat(group_regressions, axis=1)
                                quadratic_regressions.append(group_regressions)
                            quadratic_regressions = pd.concat(
                                quadratic_regressions, axis=0
                            )
                            name_tag_regressions.append(quadratic_regressions)
                        except LinAlgError as e:
                            print(f"{e}\n{traceback.format_exc()}")
                            pass
                if name_tag_regressions:
                    name_tag_regressions = pd.concat(name_tag_regressions, axis=0)

                    levels_to_remap = [
                        QuantileRegStrs.DEPENDENT_VAR,
                        QuantileRegStrs.INDEPENDENT_VARS,
                    ]
                    pretty_index = name_tag_regressions.index.set_levels(
                        [
                            prettify_index_level(
                                str_mapper,
                                QuantileRegStrs.QUADRATIC_VAR_SUFFIX,
                                level,
                                level_id,
                                levels_to_remap,
                            )
                            for level, level_id in zip(
                                name_tag_regressions.index.levels,
                                name_tag_regressions.index.names,
                            )
                        ],
                        level=name_tag_regressions.index.names,
                    )
                    name_tag_regressions.index = pretty_index

                    name_tag_regressions.to_parquet(dep_var_name_parquet)

                    name_tag_regressions.to_excel(dep_var_name_excel)
                    book = openpyxl.load_workbook(dep_var_name_excel)
                    for sheet_name in book.sheetnames:
                        sheet = book[sheet_name]
                        auto_adjust_columns_width(sheet)
                    book.save(dep_var_name_excel)
    return regressions, regressions_info


def regex_symbol_replacement(match):
    return rf"\{match.group(0)}"
